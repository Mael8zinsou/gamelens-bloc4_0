"""Generation du support de soutenance : docs/support_soutenance.md vers .pptx.

Le document Markdown est la SOURCE, le PPTX en est une sortie. Une retouche
faite dans PowerPoint est perdue a la regeneration suivante : corriger le .md.

## Le systeme visuel

Un premier jet donnait trente-trois diapositives de la meme forme, titre plus
puces, avec une pastille rouge « eliminatoire » tamponnee sur chacune. C'etait un
support concu pour la grille et non pour deux personnes assises trente minutes :
le marquage de conformite s'adresse au correcteur, pas au public.

Ce qui a change :

- **Sept formes de diapositive** au lieu d'une, choisies par le role : couverture,
  puces, figure de tete, visuel plein cadre, preuve, demonstration, duo.
- **Le marquage de conformite est descendu en pied de page**, discret. Le jury a
  la grille en main ; il n'a pas besoin qu'on la lui recite a l'ecran.
- **Un bandeau de section** en surtitre donne le rythme sans couter de temps,
  la ou des intercalaires en auraient consomme.
- **Une palette validee** plutot que du gris : bleu, orange et aqua sont les
  trois premiers emplacements categoriels de la palette de reference, valides
  ensemble en mode clair.

    python outils/generer_support.py
"""

from __future__ import annotations

import re
import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

RACINE = Path(__file__).resolve().parent.parent
SOURCE = RACINE / "docs" / "support_soutenance.md"
SORTIE = RACINE / "docs" / "support_soutenance.pptx"
CAPTURES = RACINE / "docs" / "captures"
PREUVES = RACINE / "docs" / "preuves"

LARGEUR = Inches(13.333)
HAUTEUR = Inches(7.5)
MARGE = Inches(0.72)
UTILE = LARGEUR - 2 * MARGE

HAUT_CONTENU = Inches(1.92)
BAS_CONTENU = HAUTEUR - Inches(0.72)
HAUTEUR_CONTENU = BAS_CONTENU - HAUT_CONTENU

ENCRE = RGBColor(0x0B, 0x0B, 0x0B)
ENCRE_2 = RGBColor(0x52, 0x51, 0x4E)
SURFACE = RGBColor(0xFC, 0xFC, 0xFB)
PANNEAU = RGBColor(0xF1, 0xF2, 0xF4)
BLEU = RGBColor(0x2A, 0x78, 0xD6)
ORANGE = RGBColor(0xEB, 0x68, 0x34)
AQUA = RGBColor(0x1B, 0xAF, 0x7A)
FILET = RGBColor(0xC9, 0xCE, 0xD5)
BLANC = RGBColor(0xFF, 0xFF, 0xFF)

ELIMINATOIRES = {"C4.2.1", "C4.2.2", "C4.2.3"}

REPLIS = {
    "cloisonnement_roles": ("c11_cloisonnement_roles.txt",
                            "Le meme role, refuse puis servi"),
    "airflow": ("c14_airflow_dags.txt", "Les quatre DAG et leurs derniers runs"),
    "grafana": ("c19_supervision_indicateurs.txt",
                "Les memes indicateurs, lus en SQL"),
}


# ---------------------------------------------------------------- analyse ----

def analyser(texte: str) -> list[dict]:
    diapos: list[dict] = []
    motif = re.compile(
        r"^## (D\d+)\. ([^\n]+)$\n\n((?:    \w+: [^\n]*\n)+)"
        r"(.*?)(?=^## D\d+\.|^# |\Z)",
        re.M | re.S,
    )
    for ident, titre, bloc, corps in motif.findall(texte):
        champs = dict(l.strip().split(": ", 1) for l in bloc.strip().splitlines())
        puces = re.search(r"### Puces\n(.*?)(?=\n### |\Z)", corps, re.S)
        notes = re.search(r"### Notes\n(.*?)(?=\n### |\Z)", corps, re.S)
        diapos.append({
            "ident": ident,
            "titre": titre.strip(),
            "type": champs.get("type", "puces"),
            "section": champs.get("section", ""),
            "minute": champs.get("minute", ""),
            "duree": champs.get("duree", "0:00"),
            "competence": champs.get("competence", "-"),
            "criteres": champs.get("criteres", "-"),
            "visuel": champs.get("visuel", "-"),
            "chiffre": champs.get("chiffre", ""),
            "legende": champs.get("legende", ""),
            "puces": [
                re.sub(r"\*\*(.+?)\*\*", r"\1", l.strip()[2:]).strip()
                for l in (puces.group(1).splitlines() if puces else [])
                if l.strip().startswith("- ")
            ],
            "notes": (notes.group(1).strip() if notes else ""),
        })
    return diapos


# ------------------------------------------------------------- primitives ----

def _cadre(diapo, gauche, haut, largeur, hauteur):
    zone = diapo.shapes.add_textbox(gauche, haut, largeur, hauteur)
    tf = zone.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    return tf


def _ligne(tf, contenu, taille, couleur=ENCRE, gras=False, avant=0,
           aligne=PP_ALIGN.LEFT, premier=False, police="Segoe UI", interligne=None):
    p = tf.paragraphs[0] if premier else tf.add_paragraph()
    p.alignment = aligne
    p.space_before = Pt(avant)
    if interligne:
        p.line_spacing = interligne
    r = p.add_run()
    r.text = contenu
    r.font.size = Pt(taille)
    r.font.bold = gras
    r.font.color.rgb = couleur
    r.font.name = police
    return p


def _rect(diapo, gauche, haut, largeur, hauteur, fond, bord=None, rayon=False,
          epaisseur=1.2):
    forme = diapo.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if rayon else MSO_SHAPE.RECTANGLE,
        gauche, haut, largeur, hauteur,
    )
    if fond is None:
        forme.fill.background()
    else:
        forme.fill.solid()
        forme.fill.fore_color.rgb = fond
    if bord is None:
        forme.line.fill.background()
    else:
        forme.line.color.rgb = bord
        forme.line.width = Pt(epaisseur)
    forme.shadow.inherit = False
    return forme


def _fond(diapo):
    _rect(diapo, 0, 0, LARGEUR, HAUTEUR, SURFACE)


def _image(diapo, fichier, gauche, haut, largeur, hauteur):
    from PIL import Image

    with Image.open(fichier) as im:
        rapport = im.width / im.height
    if largeur / hauteur > rapport:
        h, w = hauteur, int(hauteur * rapport)
    else:
        w, h = largeur, int(largeur / rapport)
    diapo.shapes.add_picture(str(fichier), gauche + (largeur - w) // 2,
                             haut + (hauteur - h) // 2, width=w, height=h)


def accent_de(diapo_donnees):
    return ORANGE if diapo_donnees["competence"] in ELIMINATOIRES else BLEU


# ------------------------------------------------------------- ossature ------

def entete(diapo, d):
    accent = accent_de(d)
    if d["section"]:
        tf = _cadre(diapo, MARGE, Inches(0.62), UTILE, Inches(0.3))
        _ligne(tf, d["section"].upper(), 11, accent, gras=True, premier=True)
        haut_titre = Inches(0.92)
    else:
        haut_titre = Inches(0.72)

    tf = _cadre(diapo, MARGE, haut_titre, Inches(11.0), Inches(0.9))
    _ligne(tf, d["titre"], 29, ENCRE, gras=True, premier=True)

    _rect(diapo, MARGE, Inches(1.72), Inches(1.15), Emu(28000), accent)


def pied(diapo, d, numero, total):
    accent = accent_de(d)
    _rect(diapo, MARGE, HAUTEUR - Inches(0.62), Emu(28000), Inches(0.16), accent)

    morceaux = [d["ident"], d["minute"]]
    if d["competence"] != "-":
        morceaux.append(d["competence"])
    if d["criteres"] != "-":
        morceaux.append(f"criteres {d['criteres']}")
    tf = _cadre(diapo, MARGE + Inches(0.14), HAUTEUR - Inches(0.66),
                Inches(9.0), Inches(0.3))
    _ligne(tf, "   ·   ".join(morceaux), 9, ENCRE_2, premier=True)

    tf = _cadre(diapo, LARGEUR - MARGE - Inches(1.6), HAUTEUR - Inches(0.66),
                Inches(1.6), Inches(0.3))
    _ligne(tf, f"{numero} / {total}", 9, ENCRE_2, aligne=PP_ALIGN.RIGHT, premier=True)


# ----------------------------------------------------------------- formes ----

def taille_ajustee(puces, largeur_po, hauteur_po, depart, interligne):
    taille = depart
    while taille > 10:
        par_ligne = max(12, int(largeur_po * 72 / (taille * 0.5)))
        lignes = sum(max(1, -(-len(p) // par_ligne)) for p in puces)
        h = (lignes * taille * 1.34 + (len(puces) - 1) * interligne) / 72
        if h <= hauteur_po:
            return taille
        taille -= 0.5
    return taille


def poser_puces(diapo, d, gauche, haut, largeur, hauteur, depart=17.5):
    accent = accent_de(d)
    interligne = 16
    taille = taille_ajustee(d["puces"], largeur / 914400, hauteur / 914400,
                            depart, interligne)
    for i, puce in enumerate(d["puces"]):
        y = haut + Inches(0.045) + int(i * (hauteur - Inches(0.1)) / len(d["puces"]))
        _rect(diapo, gauche, y + Inches(0.09), Inches(0.075), Inches(0.075), accent)
        tf = _cadre(diapo, gauche + Inches(0.26), y, largeur - Inches(0.26),
                    int(hauteur / len(d["puces"])))
        _ligne(tf, puce, taille, ENCRE, premier=True, interligne=1.15)


def forme_couverture(diapo, d):
    _fond(diapo)
    _rect(diapo, 0, 0, Inches(0.28), HAUTEUR, BLEU)
    tf = _cadre(diapo, Inches(1.3), Inches(2.35), Inches(10.6), Inches(1.5))
    _ligne(tf, d["titre"], 44, ENCRE, gras=True, premier=True)
    if d["legende"]:
        _ligne(tf, d["legende"], 19, ENCRE_2, avant=14)
    _rect(diapo, Inches(1.3), Inches(4.35), Inches(1.6), Emu(38000), ORANGE)
    if d["puces"]:
        tf = _cadre(diapo, Inches(1.3), Inches(4.75), Inches(10.6), Inches(1.6))
        for i, puce in enumerate(d["puces"]):
            _ligne(tf, puce, 15, ENCRE_2, premier=(i == 0), avant=0 if not i else 9)


def forme_puces(diapo, d):
    _fond(diapo)
    entete(diapo, d)
    poser_puces(diapo, d, MARGE, HAUT_CONTENU, UTILE, HAUTEUR_CONTENU)


def forme_visuel(diapo, d):
    _fond(diapo)
    entete(diapo, d)
    hauteur = HAUTEUR_CONTENU - (Inches(0.42) if d["legende"] else 0)
    chemin = RACINE / "docs" / d["visuel"]
    if chemin.exists():
        _image(diapo, chemin, MARGE, HAUT_CONTENU, UTILE, hauteur)
    else:
        _rect(diapo, MARGE, HAUT_CONTENU, UTILE, hauteur, PANNEAU, FILET, rayon=True)
    if d["legende"]:
        tf = _cadre(diapo, MARGE, BAS_CONTENU - Inches(0.34), UTILE, Inches(0.34))
        _ligne(tf, d["legende"], 13, ENCRE_2, premier=True)


def forme_chiffre(diapo, d):
    _fond(diapo)
    entete(diapo, d)
    tf = _cadre(diapo, MARGE, HAUT_CONTENU + Inches(0.3), Inches(4.7), Inches(2.2))
    _ligne(tf, d["chiffre"], 76, ENCRE, gras=True, premier=True)
    if d["legende"]:
        _ligne(tf, d["legende"], 14, ENCRE_2, avant=6)
    if d["puces"]:
        poser_puces(diapo, d, MARGE + Inches(5.3), HAUT_CONTENU,
                    UTILE - Inches(5.3), HAUTEUR_CONTENU, depart=16)


def forme_duo(diapo, d):
    """Deux colonnes : la puce qui contient ' | ' se scinde en deux."""
    _fond(diapo)
    entete(diapo, d)
    gauche_titre, droite_titre = (d["legende"].split(" | ") + ["", ""])[:2]
    colonne = (UTILE - Inches(0.5)) // 2
    for i, (titre, accent) in enumerate(((gauche_titre, ENCRE_2), (droite_titre, ORANGE))):
        x = MARGE + i * (colonne + Inches(0.5))
        tf = _cadre(diapo, x, HAUT_CONTENU, colonne, Inches(0.36))
        _ligne(tf, titre.upper(), 11.5, accent, gras=True, premier=True)
        _rect(diapo, x, HAUT_CONTENU + Inches(0.38), colonne, Emu(19000), FILET)
        items = [p.split(" | ")[i] for p in d["puces"] if " | " in p]
        faux = dict(d, puces=items)
        if items:
            poser_puces(diapo, faux, x, HAUT_CONTENU + Inches(0.62), colonne,
                        HAUTEUR_CONTENU - Inches(0.62), depart=15.5)


def forme_direct(diapo, d):
    _fond(diapo)
    entete(diapo, d)
    largeur_texte = Inches(6.0)
    poser_puces(diapo, d, MARGE, HAUT_CONTENU, largeur_texte, HAUTEUR_CONTENU)

    x = MARGE + largeur_texte + Inches(0.5)
    largeur = LARGEUR - MARGE - x
    _rect(diapo, x, HAUT_CONTENU, largeur, HAUTEUR_CONTENU, PANNEAU, ORANGE,
          rayon=True, epaisseur=1.6)
    tf = _cadre(diapo, x + Inches(0.4), HAUT_CONTENU + Inches(1.15),
                largeur - Inches(0.8), Inches(2.4))
    _ligne(tf, "EN DIRECT", 20, ORANGE, gras=True, aligne=PP_ALIGN.CENTER,
           premier=True)
    _ligne(tf, d["visuel"].split(":", 1)[1].replace("_", " "), 14, ENCRE,
           aligne=PP_ALIGN.CENTER, avant=12)
    _ligne(tf, "Repli capture a la diapositive suivante", 11.5, ENCRE_2,
           aligne=PP_ALIGN.CENTER, avant=16)
    _ligne(tf, "Sans reponse en 10 secondes, y passer sans commenter", 11.5,
           ENCRE_2, aligne=PP_ALIGN.CENTER, avant=5)


def forme_capture(diapo, d):
    _fond(diapo)
    entete(diapo, d)
    identifiant = d["visuel"].split(":", 1)[1]
    fichier = CAPTURES / f"{identifiant}.png"
    hauteur = HAUTEUR_CONTENU - (Inches(0.42) if d["legende"] else 0)
    if fichier.exists():
        _image(diapo, fichier, MARGE, HAUT_CONTENU, UTILE, hauteur)
    else:
        _rect(diapo, MARGE, HAUT_CONTENU, UTILE, hauteur, PANNEAU, FILET,
              rayon=True, epaisseur=1.4)
        tf = _cadre(diapo, MARGE + Inches(1.0), HAUT_CONTENU + hauteur // 2
                    - Inches(0.55), UTILE - Inches(2.0), Inches(1.2))
        _ligne(tf, "Capture attendue", 17, ORANGE, gras=True,
               aligne=PP_ALIGN.CENTER, premier=True)
        _ligne(tf, f"docs/captures/{identifiant}.png", 13, ENCRE_2,
               aligne=PP_ALIGN.CENTER, avant=10)
        _ligne(tf, "Deposer l'image, puis relancer le generateur", 11.5, ENCRE_2,
               aligne=PP_ALIGN.CENTER, avant=5)
    if d["legende"]:
        tf = _cadre(diapo, MARGE, BAS_CONTENU - Inches(0.34), UTILE, Inches(0.34))
        _ligne(tf, d["legende"], 13, ENCRE_2, premier=True)


def forme_preuve(diapo, d, lignes: list[str]):
    _fond(diapo)
    entete(diapo, d)
    # Le panneau epouse le contenu : un extrait de six lignes n'a pas a flotter
    # dans un cadre dimensionne pour quinze. Le coefficient est cale sur du
    # Consolas 11 points a l'interligne courant.
    hauteur = min(HAUTEUR_CONTENU, Inches(0.52) + len(lignes) * Inches(0.205))
    _rect(diapo, MARGE, HAUT_CONTENU, UTILE, hauteur, PANNEAU, FILET, rayon=True)
    tf = _cadre(diapo, MARGE + Inches(0.34), HAUT_CONTENU + Inches(0.26),
                UTILE - Inches(0.68), hauteur - Inches(0.5))
    for i, ligne in enumerate(lignes):
        _ligne(tf, ligne[:112], 11, ENCRE, premier=(i == 0), police="Consolas")
    return hauteur


def forme_preuve_fichier(diapo, d):
    """Une sortie reelle, montree telle quelle, en police fixe.

    L'en-tete de provenance des captures est retire : il documente le fichier,
    il n'apporte rien a l'ecran.
    """
    chemin = PREUVES / d["visuel"].split(":", 1)[1]
    hauteur = HAUTEUR_CONTENU
    if chemin.exists():
        lignes = [l for l in chemin.read_text(encoding="utf-8").splitlines()
                  if not l.startswith("#") and l.strip()]
        hauteur = forme_preuve(diapo, d, lignes[:15])
    else:
        _fond(diapo)
        entete(diapo, d)
        _rect(diapo, MARGE, HAUT_CONTENU, UTILE, HAUTEUR_CONTENU, PANNEAU,
              FILET, rayon=True)
    if d["legende"]:
        haut_legende = min(HAUT_CONTENU + hauteur + Inches(0.16),
                           BAS_CONTENU - Inches(0.3))
        tf = _cadre(diapo, MARGE, haut_legende, UTILE, Inches(0.3))
        _ligne(tf, d["legende"], 12, ENCRE_2, premier=True)


FORMES = {
    "couverture": forme_couverture,
    "preuve": forme_preuve_fichier,
    "puces": forme_puces,
    "visuel": forme_visuel,
    "chiffre": forme_chiffre,
    "duo": forme_duo,
    "direct": forme_direct,
    "capture": forme_capture,
}


# ------------------------------------------------------------------ sortie ---

SANS_PUCES = {"visuel", "capture", "preuve"}


def construire(presentation, d, numero, total):
    diapo = presentation.slides.add_slide(presentation.slide_layouts[6])
    FORMES.get(d["type"], forme_puces)(diapo, d)

    notes = d["notes"]
    if d["type"] in SANS_PUCES and d["puces"]:
        # Ces formes occupent tout le cadre : les puces n'y ont pas de place a
        # l'ecran, mais elles restent ce qu'il faut dire devant l'image.
        a_dire = "\n".join(f"- {p}" for p in d["puces"])
        notes = f"A dire devant l'image :\n{a_dire}\n\n{notes}".strip()
    if notes:
        diapo.notes_slide.notes_text_frame.text = notes
    if d["type"] != "couverture":
        pied(diapo, d, numero, total)
    return diapo


def construire_repli(presentation, d, numero, total):
    identifiant = d["visuel"].split(":", 1)[1]
    fichier, titre = REPLIS[identifiant]
    chemin = PREUVES / fichier
    repli = dict(d, titre=titre, ident=d["ident"] + "r", section="repli",
                 puces=[], legende=f"Capture du 01/09, docs/preuves/{fichier}")

    diapo = presentation.slides.add_slide(presentation.slide_layouts[6])
    capture = CAPTURES / f"{identifiant}.png"
    if capture.exists():
        repli["visuel"] = f"capture:{identifiant}"
        forme_capture(diapo, repli)
    elif chemin.exists():
        lignes = [l for l in chemin.read_text(encoding="utf-8").splitlines()
                  if not l.startswith("#")]
        forme_preuve(diapo, repli, [l for l in lignes if l.strip()][:15])
    else:
        repli["visuel"] = f"capture:{identifiant}"
        forme_capture(diapo, repli)
    diapo.notes_slide.notes_text_frame.text = (
        f"Repli de {d['ident']}. Ne pas commenter la panne : enchainer comme si "
        "c'etait prevu."
    )
    pied(diapo, repli, numero, total)


def main(argv: list[str] | None = None) -> int:
    if not SOURCE.exists():
        print(f"{SOURCE.relative_to(RACINE)} est absent.")
        return 1

    diapos = analyser(SOURCE.read_text(encoding="utf-8"))
    if not diapos:
        print("Aucune diapositive reconnue : le format du Markdown a-t-il change ?")
        return 1

    total = len(diapos) + sum(1 for d in diapos if d["type"] == "direct")

    presentation = Presentation()
    presentation.slide_width = LARGEUR
    presentation.slide_height = HAUTEUR

    numero = 0
    deplacees = 0
    requises: list[str] = []
    souhaitables: list[str] = []
    for d in diapos:
        numero += 1
        if d["type"] in SANS_PUCES and d["puces"]:
            deplacees += len(d["puces"])
        construire(presentation, d, numero, total)
        if d["visuel"].startswith("capture:"):
            identifiant = d["visuel"].split(":", 1)[1]
            if not (CAPTURES / f"{identifiant}.png").exists():
                requises.append(identifiant)
        if d["type"] == "direct":
            identifiant = d["visuel"].split(":", 1)[1]
            if not (CAPTURES / f"{identifiant}.png").exists():
                souhaitables.append(identifiant)
            numero += 1
            construire_repli(presentation, d, numero, total)

    try:
        presentation.save(SORTIE)
    except PermissionError:
        print(f"{SORTIE.name} est ouvert dans PowerPoint : le fermer, puis "
              "relancer. Rien n'a ete ecrit.")
        return 1

    duree = sum(int(d["duree"].split(":")[0]) * 60 + int(d["duree"].split(":")[1])
                for d in diapos)
    formes = {}
    for d in diapos:
        formes[d["type"]] = formes.get(d["type"], 0) + 1

    print(f"Ecrit : {SORTIE.relative_to(RACINE)}")
    print(f"  {len(diapos)} diapositives du Markdown, {total} dans le PPTX")
    print(f"  duree annoncee : {duree // 60}:{duree % 60:02d} sur 30:00")
    print("  formes : " + ", ".join(f"{n} {t}" for t, n in sorted(formes.items())))
    if deplacees:
        print(f"  {deplacees} puces versees dans les notes : les formes a "
              f"visuel occupent tout le cadre")
    if requises:
        print(f"  CADRES VIDES, capture requise : {', '.join(sorted(set(requises)))}")
    if souhaitables:
        print(f"  replis assures par le texte, une capture les rendrait plus "
              f"lisibles : {', '.join(sorted(set(souhaitables)))}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
